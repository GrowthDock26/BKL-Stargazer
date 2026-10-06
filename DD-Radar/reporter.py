"""Berichtserzeugung: Excel-Arbeitsmappe, Word-Kurzbericht und HTML-Dashboard."""

from __future__ import annotations

import html
import json
from datetime import datetime
from pathlib import Path

import taxonomy
from config import AMPEL_GELB, AMPEL_GRUEN, AMPEL_ROT, STATUS_FEHLT, STATUS_VORHANDEN
from gap_analysis import Ergebnis, datenraum_plan, irl
from scanner import Datei

HINWEIS_KI = (
    "Dieser Bericht wurde mit KI-Unterstützung (LOGICC) erstellt. Die Zuordnung "
    "einzelner Dokumente und die Bewertung von Auffälligkeiten sind Vorschläge und "
    "vor Verwendung im Datenraum durch die Geschäftsführung bzw. die beratende "
    "Kanzlei zu prüfen. Der Bericht ersetzt keine Rechtsberatung."
)

HINWEIS_VERTRAULICH = (
    "VERTRAULICH — enthält Angaben über den Dokumentenbestand des Unternehmens. "
    "Weitergabe nur an das Transaktionsteam."
)

FARBE = {
    AMPEL_GRUEN: "C6EFCE",
    AMPEL_GELB: "FFEB9C",
    AMPEL_ROT: "FFC7CE",
}
SCHRIFT = {
    AMPEL_GRUEN: "006100",
    AMPEL_GELB: "9C6500",
    AMPEL_ROT: "9C0006",
}
KOPF_BLAU = "1F3864"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def schreibe_excel(pfad: Path, dateien: list[Datei], erg: Ergebnis,
                   stat: dict, meta: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    kopf_fill = PatternFill("solid", fgColor=KOPF_BLAU)
    kopf_font = Font(color="FFFFFF", bold=True)

    def kopfzeile(ws, spalten: list[str], breiten: list[int]) -> None:
        ws.append(spalten)
        for i, breite in enumerate(breiten, start=1):
            ws.column_dimensions[get_column_letter(i)].width = breite
            zelle = ws.cell(row=1, column=i)
            zelle.fill = kopf_fill
            zelle.font = kopf_font
            zelle.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"

    def faerbe(ws, zeile: int, spalte: int, ampel: str) -> None:
        z = ws.cell(row=zeile, column=spalte)
        if ampel in FARBE:
            z.fill = PatternFill("solid", fgColor=FARBE[ampel])
            z.font = Font(color=SCHRIFT[ampel], bold=True)

    # -- 1 Übersicht --------------------------------------------------------
    ws = wb.active
    ws.title = "Übersicht"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 70

    ws["A1"] = "DD-Radar — Datenraum-Bereitschaftsprüfung"
    ws["A1"].font = Font(size=15, bold=True, color=KOPF_BLAU)
    ws["A2"] = HINWEIS_VERTRAULICH
    ws["A2"].font = Font(italic=True, size=9)

    zeilen = [
        ("", ""),
        ("Mandant / Unternehmen", meta.get("mandant", "—")),
        ("Analysiertes Verzeichnis", meta.get("verzeichnis", "")),
        ("Prüfprofil", f"{erg.profil} — {taxonomy.PROFILE.get(erg.profil, '')}"),
        ("Erstellt am", meta.get("zeitpunkt", "")),
        ("KI-Modell (LOGICC)", meta.get("modell", "—")),
        ("Betriebsmodus", meta.get("modus", "")),
        ("", ""),
        ("BEREITSCHAFTSGRAD", f"{erg.kennzahlen['bereitschaft_prozent']} %"),
        ("Soll-Positionen im Katalog", erg.kennzahlen["anzahl_soll"]),
        ("davon vorhanden", erg.kennzahlen["vorhanden"]),
        ("davon nur teilweise", erg.kennzahlen["teilweise"]),
        ("davon fehlend", erg.kennzahlen["fehlt"]),
        ("Kritische Positionen offen", erg.kennzahlen["kritisch_offen"]),
        ("Wichtige Positionen offen", erg.kennzahlen["hoch_offen"]),
        ("Auffälligkeiten (Red Flags)", len(erg.red_flags)),
        ("", ""),
        ("Dateien im Verzeichnis", stat["anzahl_gesamt"]),
        ("davon inhaltlich ausgewertet", stat["anzahl_analysierbar"]),
        ("Duplikate", stat["anzahl_duplikate"]),
        ("Altformate (.doc/.xls/.msg …)", stat["anzahl_altformat"]),
        ("Dateien ohne Katalogzuordnung", len(erg.nicht_zugeordnet)),
        ("Gesamtgröße", f"{stat['gesamtgroesse_mb']} MB"),
    ]
    for a, b in zeilen:
        ws.append([a, b])
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "BEREITSCHAFTSGRAD":
            ws.cell(row=r, column=1).font = Font(bold=True, size=12)
            ws.cell(row=r, column=2).font = Font(bold=True, size=12)
            break

    ws.append([])
    ws.append(["Ampel je Workstream", ""])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for kuerzel, name in taxonomy.WORKSTREAMS.items():
        if kuerzel not in erg.workstream_ampeln:
            continue
        ws.append([name, erg.workstream_ampeln[kuerzel]])
        faerbe(ws, ws.max_row, 2, erg.workstream_ampeln[kuerzel])

    ws.append([])
    ws.append(["Hinweis", HINWEIS_KI])
    ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    # -- 2 Gap-Analyse ------------------------------------------------------
    ws = wb.create_sheet("Gap-Analyse")
    kopfzeile(ws, ["ID", "Workstream", "Soll-Dokument", "Priorität", "Status", "Ampel",
                   "Belege im Datenraum", "Mängel", "Warum der Investor das braucht",
                   "Rechtsgrundlage", "Prüfpunkte"],
              [10, 34, 44, 12, 12, 8, 46, 40, 56, 22, 44])
    for b in erg.befunde:
        ws.append([b.soll.id, taxonomy.workstream_name(b.soll.workstream), b.soll.titel,
                   b.soll.prioritaet, b.status, b.ampel, b.belegliste, "; ".join(b.maengel),
                   b.soll.warum, b.soll.rechtsgrundlage, " | ".join(b.soll.pruefpunkte)])
        faerbe(ws, ws.max_row, 6, b.ampel)
    for zeile in ws.iter_rows(min_row=2):
        for z in zeile:
            z.alignment = Alignment(wrap_text=True, vertical="top")

    # -- 3 Nachforderungsliste ---------------------------------------------
    ws = wb.create_sheet("Nachforderungsliste")
    kopfzeile(ws, ["#", "ID", "Priorität", "Status", "Workstream", "Aufgabe",
                   "Warum", "Bereits vorhanden", "Mängel", "Prüfpunkte",
                   "Verantwortlich", "Frist", "Erledigt"],
              [5, 10, 12, 12, 34, 50, 56, 40, 40, 44, 18, 12, 10])
    for i, e in enumerate(irl(erg), start=1):
        ws.append([i, e["id"], e["prioritaet"], e["status"], e["workstream"], e["aufgabe"],
                   e["warum"], e["vorhandene_belege"], e["maengel"], e["pruefpunkte"], "", "", ""])
        faerbe(ws, ws.max_row, 3,
               AMPEL_ROT if e["prioritaet"] == "KRITISCH"
               else AMPEL_GELB if e["prioritaet"] == "HOCH" else AMPEL_GRUEN)
    for zeile in ws.iter_rows(min_row=2):
        for z in zeile:
            z.alignment = Alignment(wrap_text=True, vertical="top")

    # -- 4 Dokumentenindex --------------------------------------------------
    ws = wb.create_sheet("Dokumentenindex")
    kopfzeile(ws, ["Datei", "Ordner", "Dokumenttyp", "Sprechender Titel", "Katalog-IDs",
                   "Datum", "Parteien", "Unterschrieben", "Vollständig", "Entwurf",
                   "Konfidenz", "Quelle", "Größe (KB)", "Geändert", "Hinweis"],
              [46, 28, 26, 40, 16, 12, 30, 13, 12, 10, 10, 12, 11, 12, 34])
    for d in dateien:
        k = d.klassifikation or {}
        ws.append([d.name, d.ordner, k.get("dokumenttyp", ""), k.get("titel_klar", ""),
                   ", ".join(k.get("checklist_ids") or []), k.get("datum", ""),
                   ", ".join(k.get("parteien") or []), k.get("unterschrieben", ""),
                   k.get("vollstaendig", ""), k.get("entwurf", ""),
                   round(float(k.get("konfidenz") or 0), 2), d.text_quelle or "—",
                   round(d.groesse / 1024), d.geaendert,
                   d.hinweis or (f"Duplikat von {d.duplikat_von}" if d.duplikat_von else "")])
    for zeile in ws.iter_rows(min_row=2):
        for z in zeile:
            z.alignment = Alignment(wrap_text=True, vertical="top")

    # -- 5 Red Flags --------------------------------------------------------
    ws = wb.create_sheet("Auffälligkeiten")
    kopfzeile(ws, ["Schwere", "Workstream", "Thema", "Fundstelle", "Hinweis", "Dokument"],
              [12, 34, 46, 28, 70, 46])
    for f in erg.red_flags:
        ws.append([f.schwere, taxonomy.workstream_name(f.workstream), f.thema,
                   f.fundstelle, f.hinweis, f.datei])
        faerbe(ws, ws.max_row, 1,
               AMPEL_ROT if f.schwere == "HOCH" else AMPEL_GELB if f.schwere == "MITTEL" else AMPEL_GRUEN)
    for zeile in ws.iter_rows(min_row=2):
        for z in zeile:
            z.alignment = Alignment(wrap_text=True, vertical="top")

    # -- 6 Datenraum-Struktur ----------------------------------------------
    ws = wb.create_sheet("Datenraum-Struktur")
    kopfzeile(ws, ["Aktueller Pfad", "Zielordner im Datenraum", "Katalog-ID",
                   "Dateiname-Vorschlag", "Vertraulichkeit"],
              [56, 56, 12, 60, 16])
    for p in datenraum_plan(dateien, erg):
        ws.append([p["aktuell"], p["zielordner"], p["katalog_id"],
                   p["dateiname_vorschlag"], p["vertraulichkeit"]])

    # -- 7 Klärungsbedarf ---------------------------------------------------
    ws = wb.create_sheet("Klärungsbedarf")
    kopfzeile(ws, ["Art", "Datei", "Details"], [26, 60, 70])
    for d in erg.nicht_zugeordnet:
        ws.append(["Keine Katalogzuordnung", d.rel_pfad,
                   (d.klassifikation or {}).get("dokumenttyp", "") or "unbestimmt"])
    for d in erg.duplikate:
        ws.append(["Duplikat", d.rel_pfad, f"inhaltsgleich mit {d.duplikat_von}"])
    for d in erg.altformate:
        ws.append(["Altformat", d.rel_pfad, d.hinweis])

    wb.save(pfad)
    return pfad


# ---------------------------------------------------------------------------
# Word-Kurzbericht
# ---------------------------------------------------------------------------

def schreibe_docx(pfad: Path, erg: Ergebnis, stat: dict, meta: dict) -> Path:
    import docx
    from docx.shared import Pt, RGBColor

    dok = docx.Document()
    dok.add_heading("Due-Diligence-Bereitschaft: Statusbericht", level=0)

    p = dok.add_paragraph()
    lauf = p.add_run(HINWEIS_VERTRAULICH)
    lauf.italic = True
    lauf.font.size = Pt(8)

    dok.add_paragraph(
        f"Unternehmen: {meta.get('mandant', '—')}   |   "
        f"Profil: {erg.profil}   |   Stand: {meta.get('zeitpunkt', '')}"
    )

    dok.add_heading("1. Ergebnis auf einen Blick", level=1)
    k = erg.kennzahlen
    absatz = dok.add_paragraph()
    lauf = absatz.add_run(f"Bereitschaftsgrad: {k['bereitschaft_prozent']} %")
    lauf.bold = True
    lauf.font.size = Pt(14)
    lauf.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    dok.add_paragraph(
        f"Von {k['anzahl_soll']} erwarteten Datenraum-Positionen sind {k['vorhanden']} "
        f"vollständig belegt, {k['teilweise']} nur teilweise und {k['fehlt']} gar nicht. "
        f"{k['kritisch_offen']} als kritisch eingestufte Positionen sind offen. "
        f"Ausgewertet wurden {stat['anzahl_analysierbar']} von {stat['anzahl_gesamt']} Dateien "
        f"({stat['gesamtgroesse_mb']} MB)."
    )

    dok.add_heading("2. Ampel je Arbeitsstrang", level=1)
    tab = dok.add_table(rows=1, cols=3)
    tab.style = "Light Grid Accent 1"
    for i, titel in enumerate(("Arbeitsstrang", "Ampel", "Offene Positionen")):
        tab.rows[0].cells[i].text = titel
    for kuerzel, ampel in erg.workstream_ampeln.items():
        offen = [b for b in erg.befunde
                 if b.soll.workstream == kuerzel and b.status != STATUS_VORHANDEN]
        zeile = tab.add_row().cells
        zeile[0].text = taxonomy.workstream_name(kuerzel)
        zeile[1].text = ampel
        zeile[2].text = str(len(offen))

    dok.add_heading("3. Vordringliche Lücken", level=1)
    posten = [e for e in irl(erg) if e["prioritaet"] == "KRITISCH"][:12]
    if posten:
        for e in posten:
            dok.add_paragraph(f"{e['id']} — {e['aufgabe']}", style="List Number")
            unter = dok.add_paragraph(e["warum"])
            unter.runs[0].font.size = Pt(9)
            unter.runs[0].italic = True
    else:
        dok.add_paragraph("Keine kritischen Lücken festgestellt.")

    dok.add_heading("4. Auffälligkeiten in vorhandenen Dokumenten", level=1)
    hoch = [f for f in erg.red_flags if f.schwere == "HOCH" and f.datei != "—"][:12]
    if hoch:
        for f in hoch:
            dok.add_paragraph(
                f"{f.thema} ({f.fundstelle}) — {f.hinweis} [Quelle: {f.datei}]",
                style="List Bullet")
    else:
        dok.add_paragraph("Keine Auffälligkeiten mit hoher Schwere in den ausgewerteten "
                          "Dokumenten festgestellt.")

    dok.add_heading("5. Empfohlene nächste Schritte", level=1)
    for schritt in (
        "Nachforderungsliste (Excel, Blatt 'Nachforderungsliste') je Position mit "
        "Verantwortlichem und Frist versehen.",
        "Kritische Positionen zuerst schließen — sie blockieren den Datenraum-Start.",
        "Dokumente ohne Unterschrift durch unterzeichnete Endfassungen ersetzen.",
        "Datenraum nach der Struktur im Blatt 'Datenraum-Struktur' aufbauen; "
        "Duplikate und Entwürfe bleiben außen vor.",
        "Vor Freigabe: anwaltliche Durchsicht der als auffällig markierten Dokumente.",
    ):
        dok.add_paragraph(schritt, style="List Bullet")

    dok.add_heading("6. Hinweise", level=1)
    hinweis = dok.add_paragraph(HINWEIS_KI)
    hinweis.runs[0].font.size = Pt(9)
    quelle = dok.add_paragraph(
        f"Analyse: DD-Radar (BKL Legal OS) · KI-Endpunkt: LOGICC · Modell: "
        f"{meta.get('modell', '—')} · Betriebsmodus: {meta.get('modus', '')}. "
        "Alle Aufrufe sind in .dd-radar/audit-log.jsonl protokolliert."
    )
    quelle.runs[0].font.size = Pt(8)

    dok.save(pfad)
    return pfad


# ---------------------------------------------------------------------------
# HTML-Dashboard
# ---------------------------------------------------------------------------

_HTML_KOPF = """<!doctype html>
<meta charset="utf-8">
<title>DD-Radar — {mandant}</title>
<style>
 :root {{ --blau:#1F3864; --gruen:#1a7f45; --gelb:#9C6500; --rot:#9C0006; }}
 body {{ font-family:"Segoe UI",Arial,sans-serif; margin:0 auto; max-width:1100px;
        padding:28px; color:#1a1a1a; }}
 h1 {{ color:var(--blau); margin-bottom:2px; }}
 .meta {{ color:#666; font-size:13px; margin-bottom:22px; }}
 .karten {{ display:flex; gap:14px; flex-wrap:wrap; margin:18px 0 26px; }}
 .karte {{ flex:1 1 150px; border:1px solid #e0e0e0; border-radius:10px; padding:14px 16px; }}
 .karte .zahl {{ font-size:28px; font-weight:700; color:var(--blau); }}
 .karte .label {{ font-size:12px; color:#666; text-transform:uppercase; letter-spacing:.4px; }}
 table {{ border-collapse:collapse; width:100%; font-size:13px; margin-bottom:26px; }}
 th,td {{ border-bottom:1px solid #e6e6e6; padding:7px 9px; text-align:left; vertical-align:top; }}
 th {{ background:var(--blau); color:#fff; font-weight:600; }}
 .GRÜN {{ color:var(--gruen); font-weight:700; }}
 .GELB {{ color:var(--gelb); font-weight:700; }}
 .ROT {{ color:var(--rot); font-weight:700; }}
 .balken {{ height:14px; background:#eee; border-radius:7px; overflow:hidden; margin:6px 0 4px; }}
 .balken > div {{ height:100%; background:linear-gradient(90deg,#1F3864,#3a6ea5); }}
 .fuss {{ font-size:11px; color:#777; border-top:1px solid #e0e0e0; padding-top:12px; }}
 @media print {{ body {{ padding:0; }} th {{ -webkit-print-color-adjust:exact; }} }}
</style>
<h1>DD-Radar — Datenraum-Bereitschaft</h1>
<div class="meta">{mandant} · Profil {profil} · Stand {zeitpunkt} · Verzeichnis: {verzeichnis}</div>
"""


def schreibe_html(pfad: Path, erg: Ergebnis, stat: dict, meta: dict) -> Path:
    e = html.escape
    k = erg.kennzahlen
    teile = [_HTML_KOPF.format(
        mandant=e(str(meta.get("mandant", "—"))),
        profil=e(erg.profil),
        zeitpunkt=e(str(meta.get("zeitpunkt", ""))),
        verzeichnis=e(str(meta.get("verzeichnis", ""))),
    )]

    teile.append(f"""<div class="balken"><div style="width:{k['bereitschaft_prozent']}%"></div></div>
<div class="karten">
  <div class="karte"><div class="zahl">{k['bereitschaft_prozent']} %</div><div class="label">Bereitschaft</div></div>
  <div class="karte"><div class="zahl">{k['vorhanden']}/{k['anzahl_soll']}</div><div class="label">Positionen belegt</div></div>
  <div class="karte"><div class="zahl">{k['kritisch_offen']}</div><div class="label">kritisch offen</div></div>
  <div class="karte"><div class="zahl">{len(erg.red_flags)}</div><div class="label">Auffälligkeiten</div></div>
  <div class="karte"><div class="zahl">{stat['anzahl_analysierbar']}</div><div class="label">Dateien ausgewertet</div></div>
</div>""")

    teile.append("<h2>Ampel je Arbeitsstrang</h2><table><tr><th>Arbeitsstrang</th><th>Ampel</th>"
                 "<th>vorhanden</th><th>teilweise</th><th>fehlt</th></tr>")
    for kuerzel, ampel in erg.workstream_ampeln.items():
        teil = [b for b in erg.befunde if b.soll.workstream == kuerzel]
        v = sum(1 for b in teil if b.status == STATUS_VORHANDEN)
        f = sum(1 for b in teil if b.status == STATUS_FEHLT)
        teile.append(f"<tr><td>{e(taxonomy.workstream_name(kuerzel))}</td>"
                     f"<td class='{ampel}'>{ampel}</td><td>{v}</td>"
                     f"<td>{len(teil) - v - f}</td><td>{f}</td></tr>")
    teile.append("</table>")

    teile.append("<h2>Vordringliche Lücken</h2><table><tr><th>ID</th><th>Priorität</th>"
                 "<th>Aufgabe</th><th>Warum</th></tr>")
    for eintrag in [x for x in irl(erg) if x["prioritaet"] in ("KRITISCH", "HOCH")][:25]:
        klasse = "ROT" if eintrag["prioritaet"] == "KRITISCH" else "GELB"
        teile.append(f"<tr><td>{e(eintrag['id'])}</td><td class='{klasse}'>{e(eintrag['prioritaet'])}</td>"
                     f"<td>{e(eintrag['aufgabe'])}</td><td>{e(eintrag['warum'])}</td></tr>")
    teile.append("</table>")

    if erg.red_flags:
        teile.append("<h2>Auffälligkeiten</h2><table><tr><th>Schwere</th><th>Thema</th>"
                     "<th>Fundstelle</th><th>Hinweis</th><th>Dokument</th></tr>")
        for f in erg.red_flags[:40]:
            klasse = {"HOCH": "ROT", "MITTEL": "GELB"}.get(f.schwere, "GRÜN")
            teile.append(f"<tr><td class='{klasse}'>{e(f.schwere)}</td><td>{e(f.thema)}</td>"
                         f"<td>{e(f.fundstelle)}</td><td>{e(f.hinweis)}</td><td>{e(f.datei)}</td></tr>")
        teile.append("</table>")

    teile.append(f"<div class='fuss'>{e(HINWEIS_VERTRAULICH)}<br>{e(HINWEIS_KI)}<br>"
                 f"KI-Endpunkt LOGICC · Modell {e(str(meta.get('modell', '—')))} · "
                 f"Betriebsmodus {e(str(meta.get('modus', '')))}</div>")

    pfad.write_text("\n".join(teile), encoding="utf-8")
    return pfad


# ---------------------------------------------------------------------------
# Rohdaten (maschinenlesbar, für Folgeläufe und Nachweiszwecke)
# ---------------------------------------------------------------------------

def schreibe_json(pfad: Path, dateien: list[Datei], erg: Ergebnis,
                  stat: dict, meta: dict) -> Path:
    daten = {
        "meta": meta,
        "statistik": stat,
        "kennzahlen": erg.kennzahlen,
        "workstream_ampeln": erg.workstream_ampeln,
        "befunde": [
            {
                "id": b.soll.id,
                "titel": b.soll.titel,
                "workstream": b.soll.workstream,
                "prioritaet": b.soll.prioritaet,
                "status": b.status,
                "ampel": b.ampel,
                "belege": [d.rel_pfad for d in b.belege],
                "maengel": b.maengel,
            }
            for b in erg.befunde
        ],
        "nachforderungen": irl(erg),
        "red_flags": [f.__dict__ for f in erg.red_flags],
        "dokumente": [d.as_dict() for d in dateien],
    }
    pfad.write_text(json.dumps(daten, ensure_ascii=False, indent=1), encoding="utf-8")
    return pfad


def zeitstempel() -> str:
    return datetime.now().strftime("%Y-%m-%d")
