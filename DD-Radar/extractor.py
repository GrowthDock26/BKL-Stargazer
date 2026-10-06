"""Textgewinnung aus PDF, DOCX, XLSX, Text und Bildern.

Gescannte PDFs (ohne verwertbare Textschicht) und Bilddateien werden nicht
per OCR verarbeitet, sondern als Seitenbild an das Vision-Modell übergeben —
dieselbe Entscheidung wie im Bescheidprüfer, weil OCR-Schichten in der Praxis
fehlerbehaftet sind.
"""

from __future__ import annotations

import base64
import logging
from pathlib import Path

from config import (
    DOCX_EXTS,
    HEAD_CHARS,
    IMAGE_EXTS,
    MAX_CHARS_PER_DOC,
    MAX_VISION_PAGES,
    PDF_EXTS,
    TAIL_CHARS,
    TEXT_EXTS,
    VISION_DPI,
    XLSX_EXTS,
)

logger = logging.getLogger(__name__)

# Ab wie vielen Zeichen je Seite gilt ein PDF als „echtes" Text-PDF
MIN_ZEICHEN_JE_SEITE = 120


class Auszug:
    """Ergebnis der Textgewinnung."""

    __slots__ = ("text", "quelle", "bilder_b64", "seiten", "hinweis")

    def __init__(self, text: str = "", quelle: str = "", bilder_b64: list[str] | None = None,
                 seiten: int = 0, hinweis: str = "") -> None:
        self.text = text
        self.quelle = quelle
        self.bilder_b64 = bilder_b64 or []
        self.seiten = seiten
        self.hinweis = hinweis

    def __bool__(self) -> bool:
        return bool(self.text.strip() or self.bilder_b64)


def kuerze(text: str) -> str:
    """Beschneidet lange Dokumente auf Anfang + Ende (Unterschriftenseite!)."""
    text = " ".join(text.split()) if text.count("\n") > 3000 else text
    if len(text) <= MAX_CHARS_PER_DOC:
        return text
    return (
        text[:HEAD_CHARS]
        + f"\n\n[… {len(text) - HEAD_CHARS - TAIL_CHARS} Zeichen ausgelassen …]\n\n"
        + text[-TAIL_CHARS:]
    )


def extrahiere(pfad: Path, use_vision: bool = True) -> Auszug:
    """Gewinnt Text (und ggf. Seitenbilder) aus einer Datei."""
    endung = pfad.suffix.lower()
    try:
        if endung in PDF_EXTS:
            return _pdf(pfad, use_vision)
        if endung in DOCX_EXTS:
            return _docx(pfad)
        if endung in XLSX_EXTS:
            return _xlsx(pfad)
        if endung in TEXT_EXTS:
            return _text(pfad)
        if endung in IMAGE_EXTS:
            return _bild(pfad, use_vision)
    except Exception as exc:                      # defekte Dateien nicht fatal
        logger.warning("Extraktion fehlgeschlagen (%s): %s", pfad.name, exc)
        return Auszug(hinweis=f"Nicht lesbar: {exc}")
    return Auszug(hinweis="Dateityp wird nicht ausgewertet")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _pdf(pfad: Path, use_vision: bool) -> Auszug:
    import fitz  # PyMuPDF

    with fitz.open(pfad) as doc:
        seiten = doc.page_count
        teile: list[str] = []
        for i in range(min(seiten, 60)):          # Deckel gegen 1000-Seiten-Anhänge
            teile.append(doc.load_page(i).get_text("text"))
        text = "\n".join(teile).strip()

        if len(text) >= MIN_ZEICHEN_JE_SEITE * min(seiten, 3):
            return Auszug(kuerze(text), "pdf-text", seiten=seiten)

        # Kaum Text → gescannt
        if not use_vision:
            return Auszug(kuerze(text), "pdf-text-duenn", seiten=seiten,
                          hinweis="Vermutlich gescannt; Vision-Auswertung deaktiviert")

        bilder: list[str] = []
        zoom = VISION_DPI / 72.0
        matrix = fitz.Matrix(zoom, zoom)
        for i in range(min(seiten, MAX_VISION_PAGES)):
            pix = doc.load_page(i).get_pixmap(matrix=matrix, alpha=False)
            bilder.append(base64.b64encode(pix.tobytes("png")).decode("ascii"))
        return Auszug(kuerze(text), "vision", bilder, seiten=seiten,
                      hinweis="Gescanntes PDF — Auswertung über Seitenbilder")


# ---------------------------------------------------------------------------
# Office / Text / Bild
# ---------------------------------------------------------------------------

def _docx(pfad: Path) -> Auszug:
    import docx  # python-docx

    dok = docx.Document(str(pfad))
    teile = [p.text for p in dok.paragraphs if p.text.strip()]
    for tabelle in dok.tables:
        for zeile in tabelle.rows:
            zellen = [z.text.strip() for z in zeile.cells]
            if any(zellen):
                teile.append(" | ".join(zellen))
    return Auszug(kuerze("\n".join(teile)), "docx")


def _xlsx(pfad: Path) -> Auszug:
    import openpyxl

    wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    teile: list[str] = []
    try:
        for ws in wb.worksheets:
            teile.append(f"### Tabellenblatt: {ws.title}")
            for i, zeile in enumerate(ws.iter_rows(values_only=True)):
                if i >= 80:                       # Kopfbereich genügt zur Typerkennung
                    teile.append("[… weitere Zeilen ausgelassen …]")
                    break
                werte = [str(w) for w in zeile if w not in (None, "")]
                if werte:
                    teile.append(" | ".join(werte))
    finally:
        wb.close()
    return Auszug(kuerze("\n".join(teile)), "xlsx")


def _text(pfad: Path) -> Auszug:
    roh = pfad.read_text(encoding="utf-8", errors="replace")
    return Auszug(kuerze(roh), "text")


def _bild(pfad: Path, use_vision: bool) -> Auszug:
    if not use_vision:
        return Auszug(hinweis="Bilddatei — Vision-Auswertung deaktiviert")
    import fitz

    pix = fitz.Pixmap(str(pfad))
    if pix.n > 3:                                  # Alphakanal entfernen
        pix = fitz.Pixmap(fitz.csRGB, pix)
    b64 = base64.b64encode(pix.tobytes("png")).decode("ascii")
    return Auszug("", "vision", [b64], seiten=1, hinweis="Bilddatei — Auswertung über Vision")
